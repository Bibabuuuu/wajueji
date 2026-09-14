"""Build a reviewable technical draft in the official proposal template."""
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "evidence" / "submission-rules-2026-09-14" / "第十五届中国创新创业大赛AI+工程机械专业赛揭榜挂帅参赛项目计划书建议模板.docx"
OUTPUT = Path(__file__).resolve().parent / "复杂介质工况自主挖掘-项目方案技术草稿-v0.1.docx"


def set_run_font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def replace(paragraph, text, *, placeholder=False):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run)
    if placeholder:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.first_line_indent = Pt(24)


def insert_after(paragraph, text, *, size=16, bold=True, placeholder=False):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    target = paragraph._parent.add_paragraph()
    target._p.getparent().remove(target._p)
    new_p.addnext(target._p)
    target._p.getparent().remove(new_p)
    run = target.add_run(text)
    set_run_font(run, "黑体" if bold else "宋体", size, bold)
    if placeholder:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return target


def main():
    doc = Document(TEMPLATE)
    doc.core_properties.title = "面向复杂介质动态地形的约束感知与反馈连续自主挖装系统"
    doc.core_properties.subject = "揭榜命题一：复杂介质工况自主挖掘"
    doc.core_properties.author = "参赛团队（待填）"
    doc.core_properties.last_modified_by = "参赛团队（待填）"
    doc.core_properties.comments = "2026-09-14技术草稿；黄色内容待团队回填；指标须以冻结评价结果更新。"

    if "正文（技术草稿）" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("正文（技术草稿）", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(12)

    p = list(doc.paragraphs)
    project = insert_after(p[1], "面向复杂介质动态地形的约束感知与反馈连续自主挖装系统", size=18)
    insert_after(project, "技术草稿 v0.1｜黄色内容待团队回填｜2026年9月14日", size=11, bold=False, placeholder=True)

    replace(p[4], "【待团队填写】牵头学校/单位全称、联合单位、科研平台与行业方向；如为联合体，列明各方任务，并附单位同意参评证明。", placeholder=True)
    replace(p[6], "技术原型已形成感知、规划、执行和物料账本四层能力：双视角深度重建10厘米地形栅格并记录覆盖与观测年龄；只在已观测走廊内选择入口；以120 Hz闭环反馈完成多关节轨迹；用固定粒子身份追踪斗内新获取、历史余料和车厢稳定存量。")
    replace(p[7], "冻结单铲基线在390400个MPM材料点、120 Hz机构步、250次MPM迭代和原材料参数下完成16190步、134.9167模拟秒的自主挖掘、提升、回转、卸料、稳定接收和返回。离土前斗内稳定携带约62.6 kg，返回后车厢稳定净增约193.2 kg，场景越界为0。")
    replace(p[8], "候选筛选已比较10个挖法，其中4个进入高保真MPM，3个通过获取保持门槛，1个因收斗跟踪超时失败。简化地形代理因修改未知单元、挖后地形误差过大和漏检物理失败而未获部署资格，由此形成“逐铲重感知、发令前复核、失败留样”的连续作业路线。修复后的连续双铲空载已完成32938个机构步、274.4833模拟秒，两入口间距0.60 m，独立审计通过；真实连续MPM正在验证，结果须在冻结后回填。")
    replace(p[9], "【待团队填写并附证明】项目负责人及成员经历、强相关论文/专利/软著、既有科研项目、实验平台和设备条件。", placeholder=True)

    replace(p[12], "研究内容1——复杂地形感知和可信状态。融合多视角深度、本体状态与接触/载荷遥测，形成带时间戳、坐标系、覆盖度和有效性的统一状态包；地形栅格显式区分已观测与未知单元，执行前复核观测新鲜度和局部表面变化。\n\n研究内容2——未知区域约束和自主选点。沿候选进土方向构造1.35 m深、0.60 m斗宽并含0.10 m余量的扫掠走廊；任一单元未知即拒绝，历史入口最小间距0.58 m；再按表面高度、局部起伏、可达性和预计收益排序，无可行候选时输出结构化失败。\n\n研究内容3——多关节重载轨迹与反馈执行。由新入口动态生成接近、切入、深入、收斗、提升、持土、回转、对车、卸料、离车和返回轨迹；每段先检查闭链、关节范围、速度、碰撞和扫掠，再由参考跟踪器执行，阶段转换同时检查到位、速度和载荷。\n\n研究内容4——连续地形演化和物料账本。一个连续序列共用同一MPM实例，使应力、塑性历史、散落土、斗内余料和车内库存自然延续；每铲锁定新的获取群体，卸料用固定粒子在车厢核心区连续5秒且速度不超过0.05 m/s作为稳定接收代理。独立评价从完整初态重建，禁止用位置/速度快照冒充完整复位。\n\n研究内容5——失败筛选和统计评价。采用几何、空载动力学、带土获取、完整运输卸料、连续序列的分层筛选；冻结后按预先声明的场景和种子报告计划数、尝试数、成功数、失败原因、装载体积、响应时延和模拟作业效率，并将算法时延与高保真仿真墙钟分开。\n\n技术路线：多传感器观测→自体过滤与地形重建→已知/未知栅格→全扫掠走廊候选→机构几何与闭链检查→动态参考轨迹→重载反馈执行→获取/持土判定→回转卸料与稳定接收→返回→新一轮观测。离线MPM只用于少量参数模板比较，线上选点不等待分钟级物理预演。\n\n创新点：把未知区域判断扩展至完整斗宽与进退扫掠走廊；以真实挖后状态重感知代替未获准入的演化代理；用固定物料身份区分本铲新获取、历史余料和稳定车载；将几何、空载、带载、接收与连续性审计串成故障闭锁链；为后续低维参数策略或安全残差强化学习保留接口。")
    replace(p[14], "预期达到并如实报告命题四项指标：任务成功率、满斗率、响应时延和相对人工效率。当前已有完整单铲和双铲空载证据；真实双铲/三铲、固定斗容体积、全链P50/P95/P99时延、同初态人工对照和独立样本统计在9月18日前冻结口径、9月21日前锁定结果。若样本量不足以证明95%，将同时报告点估计、样本数和不确定性，不以单次演示替代统计结论。")
    replace(p[15], "形成可复现的约束选点、动态轨迹、反馈控制、MPM耦合、物料账本和独立审计源码；交付算法纯文本、环境与来源清单、运行配置、代表数据、失败样本、三维模型以及版本和文件哈希。根据评审要求整理软件著作权、论文或专利材料。")
    replace(p[16], "成果载体包括：项目方案DOCX/PDF、测试报告及指标附表、连续自主作业MP4、未剪辑原始记录索引、STEP/GLB模型、原创源码与上游来源声明、代表运行数据和标注规范。所有定量主张映射到唯一运行编号，失败样本与成功样本同等留存。")
    replace(p[17], "成果面向露天矿山、土方施工和装卸场景，可作为自主挖装控制软件、仿真验证工具链和测试规范交付。工程化阶段将用真实相机/激光雷达、编码器和液压压力替换仿真接口，并由榜单单位提供目标机型、斗容、液压接口、典型介质与人工生产率基准完成标定。")

    replace(p[19], "2026年9月14—18日完成连续三铲、体积与时延口径并冻结策略和评价集；9月19—21日完成独立评价、有限人工对照及条件允许时的介质对比；9月22日锁定结果、测试报告和方案；9月23日完成连续演示视频与文档渲染；9月24日完成源码、数据、模型复现和内部上传验收；9月25日只处理上传或格式阻断。\n\n【待团队确认正式项目周期】建议按0—3个月实机接口与传感器标定、4—6个月复杂介质数据与约束策略验证、7—9个月现场连续装车试验、10—12个月可靠性与产业化验收设置里程碑。", placeholder=True)
    replace(p[21], "系统采用传感器、规划、控制和物料状态模块化接口，高保真MPM用于碰撞、阻力和地形演化的开发验证，不能替代实土标定。工程化阶段需由徐工提供目标机型斗容、关节/液压接口、传感器布置、典型介质和人工生产率基准。产品形态包括自主挖装控制软件、仿真验证工具链及数据/测试规范，价值来自减少人工暴露、降低碰撞停机、提升循环一致性并形成可追溯评价。")
    replace(p[23], "【待团队填写】负责人、感知/规划/控制/仿真/材料负责人，周例会和问题闭环机制，源码与数据权限，知识产权归属，联合单位分工；明确技术负责人审核定量主张，材料负责人审核脱敏、附件、盖章和平台回执。", placeholder=True)
    replace(p[25], "【待团队填写】榜额、学校配套和预算依据。硬件按经批准的设备购置、改造或云算力计入；当前性能分析不支持为单序列增加同机第二张显卡。若冻结评价量超过50小时或单机可靠性威胁9月21日锁结果，优先租用第二台兼容Linux/CUDA主机并行独立序列。预算表应覆盖设备、材料测试、差旅会议、知识产权、劳务咨询及其他支出。", placeholder=True)
    risk = insert_after(p[27], "连续第二或第三铲受挖后坑壁、斗内余料或车载库存影响时，每铲重新观测，完整走廊不满足即停止并留样；MPM计算慢时先冻结可靠闭环，独立序列可分配第二台兼容Linux/CUDA主机，不拆分同一物理序列；满斗率使用固定斗容和体积算法，质量只作诊断并执行已知体积与分辨率检查；不同介质只采用来源明确的参数或列为后续实土标定；仿真作品边界、视频时长、STEP/GLB和源码入口由团队通过登录页或组委会确认；校内法人证书、知识产权证明与行政章自9月14日起并行办理。", size=12, bold=False)
    risk.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    risk.paragraph_format.line_spacing = 1.5
    risk.paragraph_format.space_after = Pt(6)
    risk.paragraph_format.first_line_indent = Pt(24)

    for table in doc.tables:
        for row in table.rows[1:]:
            for cell in row.cells[1:]:
                if not cell.text.strip():
                    run = cell.paragraphs[0].add_run("【待填】")
                    set_run_font(run, size=10)
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not footer.text.strip():
            run = footer.add_run("技术草稿 v0.1｜定量结果与团队信息待冻结后复核")
            set_run_font(run, size=9)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
